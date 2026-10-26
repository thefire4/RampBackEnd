from datetime import datetime
from pathlib import Path

import openpyxl

SAMPLE_PROJECT = {
    "style_no": "TS1001",
    "po_no": "9001",
    "customer": "SAMPLE BRAND",
    "season": "SS27",
    "size_range": "S-XL",
    "fabric_yarn": "Sample cotton twill",
    "designer": "Test Factory",
    "weight": "200 GSM",
    "description": "TEST CHORE JACKET",
}

COLORWAYS = ["NAVY", "STONE"]
SIZES = ["S", "M", "L"]

# code, point measured, tolerance, measurement per size, grade
SAMPLE_MEASUREMENTS = [
    ("A", "Front length", 0.5, (28.5, 29.5, 30.5), 1.0),
    ("B", "Chest width", 0.375, (21.0, 22.5, 24.0), 1.5),
]

# name, location, material, article no, size, quantity, uom, supplier, contact, color per colorway
SAMPLE_LINES = [
    ("Shell fabric", "Body", "100% cotton twill", "TW-100", "58 in", 2, "yd", "Acme Mills", "Jane Doe", ("NAVY", "STONE")),
    ("Button", "Center front", "4-hole horn button", "BT-20", "24L", 6, "pcs", "Button Co", "Jane Doe", ("Navy horn", "Stone horn")),
    # No supplier on purpose, so this line should be flagged as missing info
    ("Hang tag", "Left sleeve", "Kraft paper tag", "HT-1", "2 x 3 in", 1, "pcs", None, "Jane Doe", ("Kraft", "Kraft")),
]


def build_sample_bom(path: Path) -> Path:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "BILL OF MATERIALS"

    def put(row, values):
        for col, value in values.items():
            ws.cell(row=row, column=col + 1, value=value)

    p = SAMPLE_PROJECT
    put(1, {0: "TECHNICAL PACK FOR STYLE"})
    put(2, {0: "BILL OF MATERIALS"})
    put(3, {0: "STYLE NO:", 3: p["style_no"], 7: "PO NO:", 11: p["po_no"], 15: "DATE:", 18: datetime(2026, 1, 15)})
    put(4, {0: "CUSTOMER:", 3: p["customer"], 7: "SIZE RANGE:", 11: p["size_range"]})
    put(5, {0: "SEASON:", 3: p["season"], 7: "FABRIC/YARN:", 11: p["fabric_yarn"]})
    put(6, {0: "MANUFACTURER:", 4: p["designer"], 15: "WEIGHT", 18: p["weight"]})
    put(7, {0: "DESCRIPTION:", 4: p["description"]})
    put(8, {16: "COLORWAYS"})
    put(9, {
        0: "Item", 4: "LOCATION", 5: "PRODUCT DESCRIPTION", 8: "PRODUCT CODE", 9: "SIZE",
        10: "QUANTITY", 12: "UOM", 11: "SUPPLIER NAME", 13: "SUPPLIER CONTACT",
        16: COLORWAYS[0], 17: COLORWAYS[1],
    })
    for i, (name, location, material, article, size, qty, uom, supplier, contact, colors) in enumerate(SAMPLE_LINES):
        put(10 + i, {
            0: name, 4: location, 5: material, 8: article, 9: size, 10: qty, 12: uom,
            11: supplier, 13: contact, 16: colors[0], 17: colors[1],
        })

    spec = wb.create_sheet("SPEC")
    spec.cell(row=2, column=1, value="SPECIFICATION FORM")
    spec.cell(row=3, column=1, value="STYLE NO:")
    spec.cell(row=3, column=4, value=p["style_no"])
    headers = {0: "REF", 2: "POINTS MEASURED", 8: "TOL +/-", 9: SIZES[0], 10: SIZES[1], 11: SIZES[2], 12: "GRADE"}
    for col, value in headers.items():
        spec.cell(row=8, column=col + 1, value=value)
    for i, (code, name, tol, sizes, grade) in enumerate(SAMPLE_MEASUREMENTS):
        row = 9 + i
        for col, value in {0: code, 2: name, 8: tol, 9: sizes[0], 10: sizes[1], 11: sizes[2], 12: grade}.items():
            spec.cell(row=row, column=col + 1, value=value)

    wb.save(path)
    return path