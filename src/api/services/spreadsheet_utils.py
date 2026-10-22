from __future__ import annotations

import csv
from io import BytesIO, StringIO
from pathlib import Path
from typing import Any, Iterable, Iterator, Optional

import openpyxl
import pandas as pd

from src.api.lib.material_fields import (
    KNOWN_BOM_FIELDS,
    match_bom_field,
    normalize_header,
)
from src.api.services.import_errors import DocumentImportError


def load_sheet_rows(path: Path, sheet_name: Optional[str] = None) -> list[tuple[int, tuple[Any, ...]]]:
    """Return (1-based row number, row values) for a spreadsheet tab or CSV."""
    suffix = path.suffix.lower()
    if suffix == ".csv":
        return _load_csv_rows(path)
    if suffix in {".xlsx", ".xlsm", ".xltx", ".xltm"}:
        return _load_xlsx_rows(path, sheet_name)
    raise DocumentImportError(
        f"Unsupported spreadsheet format '{suffix}'. Use Excel (.xlsx, .xlsm) or CSV.",
        status_code=415,
    )


def _load_csv_rows(path: Path) -> list[tuple[int, tuple[Any, ...]]]:
    try:
        with path.open(newline="", encoding="utf-8-sig") as handle:
            reader = csv.reader(handle)
            return [(idx, tuple(row)) for idx, row in enumerate(reader, start=1)]
    except OSError as exc:
        raise DocumentImportError(f"Could not read CSV file: {exc}") from exc
    except csv.Error as exc:
        raise DocumentImportError(f"Invalid CSV file: {exc}") from exc


def _load_xlsx_rows(path: Path, sheet_name: Optional[str]) -> list[tuple[int, tuple[Any, ...]]]:
    try:
        workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    except OSError as exc:
        raise DocumentImportError(f"Could not read Excel file: {exc}") from exc
    except Exception as exc:
        raise DocumentImportError(f"Invalid or corrupt Excel file: {exc}") from exc

    try:
        if sheet_name:
            if sheet_name not in workbook.sheetnames:
                raise DocumentImportError(f"Sheet '{sheet_name}' was not found in the workbook.")
            sheet = workbook[sheet_name]
        else:
            sheet = workbook.active
        rows: list[tuple[int, tuple[Any, ...]]] = []
        for index, row in enumerate(sheet.iter_rows(values_only=True), start=1):
            rows.append((index, tuple(row)))
        return rows
    finally:
        workbook.close()


def list_workbook_sheets(path: Path) -> list[str]:
    suffix = path.suffix.lower()
    if suffix == ".csv":
        return ["Sheet1"]
    try:
        workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    except Exception as exc:
        raise DocumentImportError(f"Invalid or corrupt Excel file: {exc}") from exc
    try:
        return list(workbook.sheetnames)
    finally:
        workbook.close()


def pick_sheet_names(path: Path, keywords: Iterable[str]) -> list[str]:
    keywords = [k.lower() for k in keywords]
    names = list_workbook_sheets(path)
    matched = [name for name in names if any(k in name.lower() for k in keywords)]
    if not matched:
        return names

    def rank(name: str) -> tuple[int, int]:
        lowered = name.lower()
        score = 0
        for keyword in keywords:
            if keyword in lowered:
                score += 10 + len(keyword)
        if "bill of material" in lowered:
            score += 50
        if lowered == "bom":
            score += 40
        if "old" in lowered or "eval" in lowered or "approval" in lowered:
            score -= 30
        return score, -len(name)

    return sorted(matched, key=rank, reverse=True)


def row_has_content(row: tuple[Any, ...]) -> bool:
    return any(cell is not None and str(cell).strip() for cell in row)


def cell_text(value: Any) -> Optional[str]:
    if value is None:
        return None
    text = str(value).strip()
    return text or None


def find_header_row(
    rows: list[tuple[int, tuple[Any, ...]]],
    *,
    min_matches: int = 2,
) -> Optional[tuple[int, dict[int, str], dict[int, str]]]:
    """
    Locate the BOM header row.

    Returns (row_number, col_index->canonical_field, col_index->raw_header).
    """
    best: Optional[tuple[int, dict[int, str], dict[int, str], int]] = None
    for row_number, row in rows[:80]:
        if not row_has_content(row):
            continue
        field_map: dict[int, str] = {}
        raw_map: dict[int, str] = {}
        for col_index, cell in enumerate(row):
            raw = cell_text(cell)
            if not raw:
                continue
            raw_map[col_index] = raw
            field = match_bom_field(raw)
            if field:
                field_map[col_index] = field
        score = len(field_map)
        if score >= min_matches and (best is None or score > best[3]):
            best = (row_number, field_map, raw_map, score)
    if best is None:
        return None
    return best[0], best[1], best[2]


def parse_quantity(value: Any) -> Optional[int]:
    if value is None or str(value).strip() == "":
        return None
    if isinstance(value, bool):
        return None
    if isinstance(value, int):
        return value
    if isinstance(value, float):
        return int(value) if value == int(value) else None
    text = str(value).strip()
    try:
        number = float(text)
        return int(number) if number == int(number) else None
    except ValueError:
        return None


def dataframe_from_sheet(path: Path, sheet_name: str) -> pd.DataFrame:
    try:
        return pd.read_excel(path, sheet_name=sheet_name, header=None, engine="openpyxl")
    except Exception as exc:
        raise DocumentImportError(f"Could not read sheet '{sheet_name}': {exc}") from exc
